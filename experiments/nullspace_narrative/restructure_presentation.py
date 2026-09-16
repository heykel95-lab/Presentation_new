"""Reorder the final deck around four main null-space result slides.

OOXML parts are edited directly to retain editable equations and videos.
The matching PDF reuses original pages and redraws changed regions at the
same coordinates, with vector result plots. The source snapshot is in review/.
"""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
import json,re,posixpath
from lxml import etree as E
import fitz
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[1];FINAL=ROOT/'Final Presentation';ASSETS=FINAL/'figures_and_images';SNAP=HERE/'review'
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
metrics=json.loads((ASSETS/'sources/nullspace_narrative_analysis.json').read_text())
PKG='http://schemas.openxmlformats.org/package/2006/relationships';CT='http://schemas.openxmlformats.org/package/2006/content-types';P14='http://schemas.microsoft.com/office/powerpoint/2010/main'
def tag(s):p,n=s.split(':');return '{'+NS[p]+'}'+n
def parse(b):return E.fromstring(b)
def dump(r):return E.tostring(r,xml_declaration=True,encoding='UTF-8',standalone=True)
def relpath(path):return posixpath.dirname(path)+'/_rels/'+posixpath.basename(path)+'.rels'
def resolve(path,target):return posixpath.normpath(posixpath.dirname(path)+'/'+target)
with ZipFile(SNAP/'before.pptx') as z:parts={n:z.read(n) for n in z.namelist()}
original=dict(parts);pres=parse(parts['ppt/presentation.xml']);pr=parse(parts['ppt/_rels/presentation.xml.rels']);cts=parse(parts['[Content_Types].xml']);targets={r.get('Id'):r.get('Target') for r in pr}
slides=[]
for i,node in enumerate(pres.find('p:sldIdLst',NS)):
 path=resolve('ppt/presentation.xml',targets[node.get(tag('r:id'))]);slides.append({'path':path,'rid':node.get(tag('r:id')),'sid':node.get('id'),'source_index':i,'new':False})
assert len(slides)==30
maxpart=max(int(re.search(r'slide(\d+)\.xml$',s['path']).group(1)) for s in slides)
maxnote=max(int(re.search(r'notesSlide(\d+)\.xml$',n).group(1)) for n in parts if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',n))
maxsid=max(int(s['sid']) for s in slides)

def clone(source,offset):
 path=f'ppt/slides/slide{maxpart+offset}.xml';note=f'ppt/notesSlides/notesSlide{maxnote+offset}.xml'
 src=slides[source-1];parts[path]=parts[src['path']];rels=parse(parts[relpath(src['path'])])
 nr=next(r for r in rels if r.get('Type').endswith('/notesSlide'));oldnote=resolve(src['path'],nr.get('Target'));nr.set('Target','../notesSlides/'+posixpath.basename(note))
 parts[relpath(path)]=dump(rels);parts[note]=parts[oldnote];nrels=parse(parts[relpath(oldnote)])
 next(r for r in nrels if r.get('Type').endswith('/slide')).set('Target','../slides/'+posixpath.basename(path));parts[relpath(note)]=dump(nrels)
 for part,typ in [(path,'slide'),(note,'notesSlide')]:E.SubElement(cts,'{'+CT+'}Override',PartName='/'+part,ContentType='application/vnd.openxmlformats-officedocument.presentationml.'+typ+'+xml')
 rid='rIdNarrative'+str(offset);E.SubElement(pr,'{'+PKG+'}Relationship',Id=rid,Type=NS['r']+'/slide',Target='slides/'+posixpath.basename(path))
 return {'path':path,'rid':rid,'sid':str(maxsid+offset),'source_index':source-1,'new':True}
newslides=[clone(21,1),clone(26,2),clone(25,3),clone(21,4)]
backup_order=[29,30,21,26,25,23,24,22,27]
ordered=slides[:20]+newslides+[slides[27]]+[slides[n-1] for n in backup_order]
assert len(ordered)==34 and len({s['sid'] for s in ordered})==34
sourcepdf=fitz.open(SNAP/'before.pdf');pdf=fitz.open()
for s in ordered:pdf.insert_pdf(sourcepdf,from_page=s['source_index'],to_page=s['source_index'])
FONT='/home/hm-panda/Desktop/usr/share/gazebo-11/media/fonts/arial.ttf';font=fitz.Font(fontfile=FONT)
boldxref=next(f[0] for f in sourcepdf[20].get_fonts() if 'Arial-Bold' in f[3])
boldbuffer=sourcepdf.extract_font(boldxref)[3];boldfont=fitz.Font(fontbuffer=boldbuffer)
NAVY=(23/255,54/255,93/255);BLACK=(20/255,)*3

def getshape(root,name):return next(x for x in root.findall('.//p:sp',NS) if x.find('p:nvSpPr/p:cNvPr',NS).get('name')==name)
def rect(sh):
 off=sh.find('p:spPr/a:xfrm/a:off',NS);ext=sh.find('p:spPr/a:xfrm/a:ext',NS);x=int(off.get('x'))/12700;y=int(off.get('y'))/12700
 return fitz.Rect(x,y,x+int(ext.get('cx'))/12700,y+int(ext.get('cy'))/12700)
def move(sh,box):
 off=sh.find('p:spPr/a:xfrm/a:off',NS);ext=sh.find('p:spPr/a:xfrm/a:ext',NS)
 for k,v in [('x',box.x0),('y',box.y0)]:off.set(k,str(round(v*12700)))
 for k,v in [('cx',box.width),('cy',box.height)]:ext.set(k,str(round(v*12700)))
def settext(sh,text,size=None):
 body=sh.find('p:txBody',NS);p=deepcopy(body.find('a:p',NS));rp=p.find('a:r/a:rPr',NS);rp=deepcopy(rp) if rp is not None else E.Element(tag('a:rPr'),sz='2000')
 if size is not None:rp.set('sz',str(round(size*100)))
 for n in list(p):
  if n.tag!=tag('a:pPr'):p.remove(n)
 for i,line in enumerate(text.split('\n')):
  if i:E.SubElement(p,tag('a:br'))
  r=E.SubElement(p,tag('a:r'));r.append(deepcopy(rp));E.SubElement(r,tag('a:t')).text=line
 for n in body.findall('a:p',NS):body.remove(n)
 body.append(p)
 return int(rp.get('sz','2000'))/100

def clear(page,box,images=0,graphics=0):
 page.add_redact_annot(box,fill=(1,1,1));page.apply_redactions(images=images,graphics=graphics)
def drawtext(page,box,text,size,colour=BLACK,bullet=False,align='left',bold=False):
 face=boldfont if bold else font;fname='ArialBoldNarrative' if bold else 'ArialNarrative'
 if bold:
  assert all(face.has_glyph(ord(c)) for c in text if not c.isspace()),text
  page.insert_font(fontname=fname,fontbuffer=boldbuffer)
 else:page.insert_font(fontname=fname,fontfile=FONT)
 x=box.x0+(20 if bullet else 0);y=box.y0+size
 if bullet:page.insert_text((box.x0,y),'•',fontname='ArialNarrative',fontsize=size*.8,color=NAVY)
 for line in text.split('\n'):
  length=face.text_length(line,fontsize=size)
  assert length<=box.width-(20 if bullet else 0)+.5,(text,'text exceeds shape width',length,box.width)
  xx=(box.x0+box.x1-length)/2 if align=='center' else box.x1-length if align=='right' else x
  page.insert_text((xx,y),line,fontname=fname,fontsize=size,color=colour);y+=size*1.1

def edit(page,root,name,text,box=None,bullet=None,size=None,colour=None):
 sh=getshape(root,name);before=rect(sh);clear(page,before)
 if box is not None:move(sh,box)
 box=rect(sh);actual=settext(sh,text,size)
 if bullet is None:bullet=sh.find('.//a:buChar',NS) is not None
 if colour is None:colour=NAVY if name in ['Slide title','Settings heading'] else BLACK
 drawtext(page,box,text,actual,colour,bullet,bold=name=='Slide title')
 return sh

def replace_plot(page,root,entry,asset,box):
 pic=next(x for x in root.findall('.//p:pic',NS) if x.find('p:nvPicPr/p:cNvPr',NS).get('name').startswith('Figure - '))
 clear(page,fitz.Rect(55,52,910,422),images=2,graphics=1)
 rp=relpath(entry['path']);rels=parse(parts[rp]);rid=pic.find('p:blipFill/a:blip',NS).get(tag('r:embed'))
 rel=next(r for r in rels if r.get('Id')==rid);rel.set('Target','../media/'+asset+'.png')
 parts['ppt/media/'+asset+'.png']=(ASSETS/(asset+'.png')).read_bytes();parts[rp]=dump(rels)
 pic.find('p:nvPicPr/p:cNvPr',NS).set('name','Figure - '+asset);move(pic,box)
 fig=fitz.open(ASSETS/(asset+'.pdf'));page.show_pdf_page(box,fig,0);fig.close()

spoken={
20:["The desired position and orientation stay fixed while joint torques apply the equivalent of a virtual 20-newton point force on link three.","First compare four conditions: no null-space torque, damping alone, conditioning alone, and both together. Conditioning is fixed at 2 newton metres and damping at 2 newton metre seconds per radian whenever each term is enabled.","Each setting has three trials. The selected terms remain active during settling and the disturbance, and all results use the same recorded 5-to-9-second interval, displayed from zero to four seconds.","After separating the two contributions, we compare conditioning magnitudes of 1.5 and 2 newton metres."],
21:["Cumulative joint motion adds projected movement in both directions. Damping alone reduces the mean from 7.60 to 5.69 degrees.","At a conditioning torque of 2 newton metres, conditioning alone gives 1.69 degrees. Combining it with damping gives 0.83 degrees, a reduction of 50.7 percent.","The curves show three-trial means and one sample standard deviation. Small net motion can still coexist with substantial movement in both directions."],
22:["The minimum singular value describes local Jacobian conditioning under the same convention in all four settings.","It decreases without null-space torque and with damping alone. Conditioning alone and the combined setting keep it nearly constant during the disturbance.","Adding damping reduced cumulative motion while retaining this nearly constant conditioning indicator. The small difference between their absolute levels does not establish a meaningful conditioning advantage."],
23:["These curves show the three-trial mean of the measured joint-one angle relative to disturbance onset. Shading is one sample standard deviation.","At the 2-newton-metre conditioning setting, repeated reversals are visible in the first-second close-up. They remain with damping, although cumulative projected motion over the full interval is lower.","Both conditioning alone and combined control end with nearly zero mean net projected motion. That cancellation does not mean the arm remained stationary."],
24:["The previous slides used 2 newton metres to make the contrasting responses visible. The lower tested conditioning torque of 1.5 newton metres produced less cumulative motion in both conditions.","Conditioning alone gives 0.29 and 1.69 degrees at the two magnitudes. With the same damping coefficient of 2 newton metre seconds per radian, the combined values are 0.19 and 0.83 degrees.","These are three-trial means. The complete curves and their sample standard deviations are retained in the backup slides.","Increasing conditioning torque did not reduce motion in these trials. Its magnitude and damping should be considered together."],
25:["Contact reduced angular error for both larger entry tilts while the desired orientation stayed fixed. The estimate retains calibration and tool-mount uncertainty.","Higher rotational stiffness increased angular error. CoC displacement changes the angular error and can carry it through zero.","Conditioning countered the tested disturbance, and adding damping reduced cumulative motion. The lower tested conditioning magnitude produced less motion in both conditions, while mean net motion remained near zero.","Future work is a more rigid tool mount and adaptive CoC selection, returning to the TCP after alignment."],
32:["This backup uses repetition one for each of the six settings, selected by the same rule throughout.","The measured joint-one angle is relative to disturbance onset. The first-second close-up compares the two conditioning magnitudes with both combined settings.","Individual reversals remain visible. The six-setting three-trial mean is available in the preceding backup comparison."]}
new_titles={21:'Cumulative joint motion',22:'Jacobian conditioning',23:'Joint motion over time',24:'Effect of conditioning torque',28:'Cumulative joint motion: all settings',29:'Jacobian conditioning: all settings',30:'Joint motion: mean of three trials',31:'Joint motion: all individual trials'}
findings={
21:('Damping alone reduced cumulative motion from 7.60° to 5.69°.','At 2 N m, adding damping reduced cumulative motion from 1.69° to 0.83°.'),
22:('The indicator decreased with no torque and damping alone.','Conditioning and combined control kept the indicator nearly constant.'),
23:('Conditioning at 2 N m shows repeated joint-motion reversals.','Adding damping reduces total motion, while reversals remain.'),
24:('At 1.5 N m, both settings produced less cumulative motion.','Adding damping reduced motion at both tested torque magnitudes.')}
notes_for_entry={};titles_for_entry={}

for n,entry in enumerate(ordered,1):
 root=parse(parts[entry['path']]);page=pdf[n-1]
 title=''.join(getshape(root,'Slide title').xpath('.//a:t/text()',namespaces=NS)) if n>1 else 'Master Thesis Presentation'
 if n in new_titles:
  title=new_titles[n];sh=getshape(root,'Slide title');box=rect(sh)
  if font.text_length(title,fontsize=22)>box.width:box=fitz.Rect(box.x0,box.y0,921.6,box.y1)
  edit(page,root,'Slide title',title,box=box,bullet=False)
 if n==20:
  # Reuse the existing four text shapes in row-wise narrative order.
  clear(page,fitz.Rect(45,353,915,472),images=0,graphics=0)
  edit(page,root,'Settings heading','Four conditions with three trials each')
  labels=[('Setting no torque','No null-space torque',fitz.Rect(60,365,450,399)),('Setting damping','Damping: 2 N m s/rad',fitz.Rect(510,365,900,399)),('Setting conditioning','Conditioning: 2 N m',fitz.Rect(60,408,450,442)),('Setting combined','Combined: conditioning + damping\n2 N m and 2 N m s/rad',fitz.Rect(510,408,900,468))]
  # Shapes are changed before drawing so old overlapping bounds do not erase new text.
  for name,text,box in labels:
   sh=getshape(root,name);move(sh,box);size=settext(sh,text);drawtext(page,box,text,size,bullet=True)
 if n in [21,22,23]:
  asset={21:'nullspace_cumulative_main',22:'nullspace_conditioning_main',23:'joint_motion_mean_main'}[n]
  box=fitz.Rect(110,52,850,422) if n==23 else fitz.Rect(100,62,860,62+760*4.2/9)
  replace_plot(page,root,entry,asset,box)
 if n==24:
  # Native editable table, with mean-only values explicitly identified.
  clear(page,fitz.Rect(55,55,910,422),images=2,graphics=1)
  tree=root.find('p:cSld/p:spTree',NS)
  for pic in list(root.findall('.//p:pic',NS)):
   if pic.find('p:nvPicPr/p:cNvPr',NS).get('name').startswith('Figure - '):pic.getparent().remove(pic)
  temp=Presentation();temp.slide_width=Pt(960);temp.slide_height=Pt(540);ts=temp.slides.add_slide(temp.slide_layouts[6])
  heading=ts.shapes.add_textbox(Pt(60),Pt(82),Pt(840),Pt(34));heading.name='Comparison quantity'
  p=heading.text_frame.paragraphs[0];p.text='Cumulative joint motion · Three-trial means';p.font.name='Arial';p.font.size=Pt(22);p.font.color.rgb=RGBColor(23,54,93)
  pp=p._p.get_or_add_pPr();pp.set('marL',str(20*12700));pp.set('indent',str(-20*12700))
  for node in [E.Element(tag('a:buSzPct'),val='80000'),E.Element(tag('a:buFont'),typeface='Arial'),E.Element(tag('a:buChar'),char='•')]:pp.insert(len(pp)-1,node)
  heading.text_frame.margin_left=heading.text_frame.margin_top=heading.text_frame.margin_right=heading.text_frame.margin_bottom=0
  frame=ts.shapes.add_table(3,3,Pt(80),Pt(147),Pt(800),Pt(174));frame.name='Conditioning torque comparison';tab=frame.table
  widths=[250,275,275]
  for c,w in enumerate(widths):tab.columns[c].width=Pt(w)
  values=[['Conditioning torque\n[N m]','Conditioning alone\n[°]','Combined\n[°]']]+[[f"{r['k_sigma_Nm']:g}",f"{r['conditioning_mean_deg']:.2f}",f"{r['combined_mean_deg']:.2f}"] for r in metrics['parameter_comparison']]
  for ri,row in enumerate(values):
   tab.rows[ri].height=Pt(58)
   for ci,text in enumerate(row):
    cell=tab.cell(ri,ci);cell.text=text;cell.vertical_anchor=MSO_ANCHOR.MIDDLE;cell.margin_top=cell.margin_bottom=Pt(3);cell.margin_left=cell.margin_right=Pt(6);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor(23,54,93) if ri==0 else RGBColor(237,242,248) if ri==1 else RGBColor(247,249,252)
    for p in cell.text_frame.paragraphs:
     p.alignment=PP_ALIGN.CENTER;p.font.name='Arial';p.font.size=Pt(20 if ri==0 else 24);p.font.bold=False;p.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(20,20,20)
  caption=ts.shapes.add_textbox(Pt(80),Pt(346),Pt(800),Pt(40));caption.name='Comparison damping coefficient'
  p=caption.text_frame.paragraphs[0];p.text='Combined damping: 2 N m s/rad';p.font.name='Arial';p.font.size=Pt(20);p.font.color.rgb=RGBColor(20,20,20);p.alignment=PP_ALIGN.CENTER
  caption.text_frame.margin_left=caption.text_frame.margin_top=caption.text_frame.margin_right=caption.text_frame.margin_bottom=0
  maxid=max(int(x.get('id')) for x in root.xpath('//*[local-name()="cNvPr"]'))
  for ix,sh in enumerate([heading,frame,caption],1):
   elem=deepcopy(sh._element);elem.xpath('.//*[local-name()="cNvPr"]')[0].set('id',str(maxid+ix));tree.append(elem)
  drawtext(page,fitz.Rect(60,82,900,116),'Cumulative joint motion · Three-trial means',22,NAVY,bullet=True)
  x=80
  for ci,w in enumerate(widths):
   for ri,row in enumerate(values):
    box=fitz.Rect(x,147+ri*58,x+w,147+(ri+1)*58);fill=NAVY if ri==0 else (237/255,242/255,248/255) if ri==1 else (247/255,249/255,252/255)
    page.draw_rect(box,color=(1,1,1),fill=fill,width=1)
    size=20 if ri==0 else 24;lines=row[ci].split('\n');textheight=size*len(lines)*1.1;tb=fitz.Rect(box.x0+6,box.y0+(58-textheight)/2,box.x1-6,box.y1)
    drawtext(page,tb,row[ci],size,(1,1,1) if ri==0 else BLACK,align='center')
   x+=w
  drawtext(page,fitz.Rect(80,346,880,386),'Combined damping: 2 N m s/rad',20,align='center')
 if n in findings:
  for j,text in enumerate(findings[n],1):edit(page,root,f'Null-space finding {j}',text)
 # Number main slides 1..24 after title, and backups B1..B9.
 if n>1:
  footer=getshape(root,'TextBox 6');label=str(n-1) if n<=25 else 'B'+str(n-25)
  if ''.join(footer.xpath('.//a:t/text()',namespaces=NS))!=label:
   oldbox=rect(footer);clear(page,fitz.Rect(900,500,924,521));size=settext(footer,label)
   box=fitz.Rect(900,oldbox.y0,921.7,oldbox.y1);move(footer,box)
   para=footer.find('p:txBody/a:p/a:pPr',NS)
   if para is None:para=E.Element(tag('a:pPr'));footer.find('p:txBody/a:p',NS).insert(0,para)
   para.set('algn','r');drawtext(page,box,label,size,(.4,.4,.4),align='right')
 if n>25:root.set('show','0')
 else:root.attrib.pop('show',None)
 parts[entry['path']]=dump(root);titles_for_entry[entry['sid']]=title
 # Preserve technical source blocks, revise spoken text where the narrative changes.
 rels=parse(parts[relpath(entry['path'])]);nr=next(r for r in rels if r.get('Type').endswith('/notesSlide'));note=resolve(entry['path'],nr.get('Target'));nroot=parse(parts[note])
 body=next(s.find('p:txBody',NS) for s in nroot.findall('.//p:sp',NS) if s.find('p:nvSpPr/p:nvPr/p:ph',NS) is not None and s.find('p:nvSpPr/p:nvPr/p:ph',NS).get('type')=='body')
 oldtext='\n'.join(body.xpath('.//a:t/text()',namespaces=NS));source='[Sources]'+oldtext.split('[Sources]',1)[1] if '[Sources]' in oldtext else '[Sources]'
 oldspoken=oldtext.split('[Sources]',1)[0].strip().splitlines();bullets=spoken.get(n,oldspoken)
 if n==28:bullets=['This backup retains all six settings with three trials each. Cumulative motion includes movement in both directions.','The complete curves use the same interval and sample-standard-deviation bands as the main comparison.']
 if n==29:bullets=['This backup compares absolute minimum singular value for all six settings under the same Jacobian convention.','The conditioning and combined settings keep the indicator nearly constant. Small absolute differences do not establish a meaningful advantage.']
 if n==31:bullets=['All three measured joint-one histories are shown for each of the six settings. The first-second close-up retains the original recorded samples.','These individual trials show variation and reversals that can be reduced in a mean curve.']
 if n==30:bullets=['This backup extends the main joint-one mean plot to all six settings. Shading is one sample standard deviation across three onset-relative histories.','Means use exact common measured times without interpolation or smoothing.']
 if n in spoken or n in [28,29,30,31]:
  for p in body.findall('a:p',NS):body.remove(p)
  for line in bullets+[source,'[Narrative]\nMain comparison: no torque, damping 2, conditioning 2, combined 2+2. Parameter comparison: conditioning 1.5 and 2, each with and without damping 2. Original 18 measured trials retained. Main plots use the same processing as six-setting backups. Combined trials were acquired in a later session, so session effects are not independently controlled.\nfigures_and_images/sources/make_nullspace_narrative.py']:
   p=E.SubElement(body,tag('a:p'));r=E.SubElement(p,tag('a:r'));E.SubElement(r,tag('a:t')).text=line
 for shp in nroot.findall('.//p:sp',NS):
  ph=shp.find('p:nvSpPr/p:nvPr/p:ph',NS)
  if ph is not None and ph.get('type')=='sldNum':
   for txt in shp.findall('.//a:fld/a:t',NS):txt.text=str(n)
 parts[note]=dump(nroot)
 notes_for_entry[entry['sid']]=bullets

# Rebuild physical slide order and section membership.
lst=pres.find('p:sldIdLst',NS)
for node in list(lst):lst.remove(node)
for entry in ordered:E.SubElement(lst,tag('p:sldId'),id=entry['sid'],attrib={tag('r:id'):entry['rid']})
sections=pres.xpath('//*[local-name()="sectionLst"]/*[local-name()="section"]')
newmembers={'Null-space control and experiments':ordered[18:24],'Conclusion and future work':ordered[24:25],'Backup':ordered[25:]}
for section in sections:
 if section.get('name') in newmembers:
  child=section.find('{'+P14+'}sldIdLst')
  for node in list(child):child.remove(node)
  for entry in newmembers[section.get('name')]:E.SubElement(child,'{'+P14+'}sldId',id=entry['sid'])
section_ids=[x.get('id') for section in sections for x in section.findall('.//{'+P14+'}sldId')]
assert section_ids==[e['sid'] for e in ordered]
parts['ppt/presentation.xml']=dump(pres);parts['ppt/_rels/presentation.xml.rels']=dump(pr);parts['[Content_Types].xml']=dump(cts)

# Portable counterpart of update-overview.ps1: regenerate its six chapter rows.
chapters=[s.get('name') for s in sections if s.get('name')!='Backup'];overview=parse(parts[ordered[3]['path']]);tree=overview.find('p:cSld/p:spTree',NS)
items=[s for s in tree.findall('p:sp',NS) if s.find('p:nvSpPr/p:cNvPr',NS).get('name').startswith('Overview item ')];template=deepcopy(items[0])
for s in list(tree.findall('p:sp',NS)):
 if s.find('p:nvSpPr/p:cNvPr',NS).get('name').startswith(('Overview item ','Overview number ')):tree.remove(s)
for row,title in enumerate(chapters):
 sh=deepcopy(template);meta=sh.find('p:nvSpPr/p:cNvPr',NS);meta.set('id',str(100+2*row));meta.set('name','Overview item '+str(row+1));move(sh,fitz.Rect(126.4,100+row*58,856.8,130+row*58));settext(sh,title,22)
 para=sh.find('p:txBody/a:p/a:pPr',NS);para.set('marL','0');para.set('indent','0');para.set('algn','l')
 for node in list(para):
  if node.tag in [tag('a:buChar'),tag('a:buAutoNum'),tag('a:buNone')]:para.remove(node)
 E.SubElement(para,tag('a:buNone'))
 num=deepcopy(sh);meta=num.find('p:nvSpPr/p:cNvPr',NS);meta.set('id',str(101+2*row));meta.set('name','Overview number '+str(row+1));move(num,fitz.Rect(72.4,100+row*58,116.4,130+row*58));settext(num,str(row+1)+'.',22);num.find('p:txBody/a:p/a:pPr',NS).set('algn','r');tree.append(num);tree.append(sh)
parts[ordered[3]['path']]=dump(overview)
# Metadata and equation catalog follow physical slide positions.
app=parse(parts['docProps/app.xml'])
for key,value in [('Slides','34'),('Notes','34'),('HiddenSlides','9')]:
 found=app.xpath('//*[local-name()="'+key+'"]')
 if found:found[0].text=value
vt='http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'
heading=app.xpath('//*[local-name()="HeadingPairs"]/*')[0]
heading[-1].find('{'+vt+'}i4').text='34'
titles=app.xpath('//*[local-name()="TitlesOfParts"]/*')[0]
for child in list(titles)[3:]:titles.remove(child)
for entry in ordered:E.SubElement(titles,'{'+vt+'}lpstr').text=titles_for_entry[entry['sid']]
titles.set('size',str(len(titles)))
parts['docProps/app.xml']=dump(app)
cat=json.loads((SNAP/'before_equations.json').read_text());by_source={e['source_index']+1:i+1 for i,e in enumerate(ordered) if not e['new']}
for item in cat:item['slide']=by_source[item['slide']]
# Verify every original equation and all video payloads before saving.
for entry in slides:
 before=parse(original[entry['path']]);after=parse(parts[entry['path']]);extract=lambda r:[E.tostring(n) for n in r.xpath('//*[local-name()="oMath"]')]
 assert extract(before)==extract(after),entry['path']
for name in original:
 if name.endswith(('.mp4','.wmv')):assert original[name]==parts[name]
with ZipFile(FINAL/'Thesis_Defense_gg0_v3.pptx','w',ZIP_DEFLATED) as z:
 for name,data in parts.items():z.writestr(name,data)
pdf.save(FINAL/'Thesis_Defense_gg0_v3.pdf',garbage=4,deflate=True)
backpdf=fitz.open();backpdf.insert_pdf(pdf,from_page=25,to_page=33);backpdf.save(FINAL/'Supplementary_slides.pdf',garbage=4,deflate=True);backpdf.close()
(ASSETS/'native_equations.json').write_text(json.dumps(cat,indent=2,ensure_ascii=False)+'\n')

# Reorder all speaking sections, preserving equations in untouched sections.
oldscript=(SNAP/'before_script.tex').read_text();preamble=oldscript.split('\\section*{',1)[0]
chunks=re.findall(r'\\section\*\{([^}]+)\}(.*?)(?=\\section\*\{|\\end\{document\})',oldscript,re.S);assert len(chunks)==30
sections_out=[];textnotes=[]
for n,entry in enumerate(ordered,1):
 title=titles_for_entry[entry['sid']];bullets=notes_for_entry[entry['sid']]
 if n in spoken or n in [28,29,30,31]:body='\n\\begin{itemize}\n'+'\n'.join('\\item '+t.replace('%','\\%') for t in bullets)+'\n\\end{itemize}\n\n'
 else:body=chunks[entry['source_index']][1]
 sections_out.append('\\section*{'+title+'}'+body);textnotes.extend([title,'-'*len(title),'\n'.join(bullets),''])
preamble=preamble.replace('all 30 slides','all 34 slides')
(FINAL/'Speaking/Thesis_Defense_Speaking_Script.tex').write_text(preamble+''.join(sections_out)+'\\end{document}\n')
(FINAL/'Speaker_notes_and_timing.txt').write_text('\n'.join(textnotes).rstrip()+'\n')
manifest=json.loads((ASSETS/'manifest.json').read_text());newassets=['nullspace_cumulative_main','nullspace_conditioning_main','joint_motion_mean_main']
manifest=[m for m in manifest if m.get('asset') not in newassets]
for asset in newassets:manifest.append({'asset':asset,'source':'sources/make_nullspace_narrative.py; sources/combined_nullspace_provenance.json','changes':'Main narrative subset: no torque, damping 2, conditioning 2, combined 2+2. Same recorded samples, means, sample SD and 5–9 s interval. The full six-setting figure is retained in hidden backup.'})
(ASSETS/'manifest.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False)+'\n')
report={'main_slides':25,'hidden_backups':9,'total_slides':34,'main_result_slides':[21,22,23,24],'backup_source_slides':backup_order,'slide_order':[{'physical':i+1,'title':titles_for_entry[e['sid']],'footer':str(i) if 0<i<25 else ('B'+str(i-24) if i>=25 else ''),'source_physical':e['source_index']+1,'hidden':i>=25,'path':e['path']} for i,e in enumerate(ordered)],'native_equations_preserved':True,'embedded_videos_preserved':True,'overview_chapters':chapters,'overview_update':'Portable XML equivalent of update-overview.ps1; same six chapter rows.','pdf_method':'Source PowerPoint PDF pages reordered, matching text/table regions redrawn and vector plots inserted.','equation_catalog_positions_updated':True}
(HERE/'validation.json').write_text(json.dumps(report,indent=2)+'\n')
print('Built 25 main slides and 9 hidden backups. Full PDF: 34 pages. Supplementary PDF: 9 pages.')
