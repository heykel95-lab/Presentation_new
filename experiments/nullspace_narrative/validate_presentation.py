"""Check narrative packaging, numerical consistency and preserved slide content."""
from pathlib import Path
from zipfile import ZipFile
from urllib.parse import unquote
import hashlib
import json
import posixpath
import re
import fitz
from lxml import etree as E
from pptx import Presentation

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
FINAL=ROOT/'Final Presentation'
ASSETS=FINAL/'figures_and_images'
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
    'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
R='{'+NS['r']+'}'
report=json.loads((HERE/'validation.json').read_text())
metrics=json.loads((ASSETS/'sources/nullspace_narrative_analysis.json').read_text())
analysis=json.loads((ASSETS/'sources/combined_nullspace_analysis.json').read_text())
with ZipFile(FINAL/'Thesis_Defense_gg0_v3.pptx') as z:
    assert z.testzip() is None
    parts={n:z.read(n) for n in z.namelist()}
with ZipFile(HERE/'review/before.pptx') as z:
    before={n:z.read(n) for n in z.namelist()}

def resolve(owner,target):
    return posixpath.normpath(posixpath.join(posixpath.dirname(owner),unquote(target))).lstrip('/')

def relpath(owner):
    return posixpath.dirname(owner)+'/_rels/'+posixpath.basename(owner)+'.rels'

# Every internal package reference must resolve after cloning and reordering.
relationships=0
for name,data in parts.items():
    if not name.endswith('.rels'):
        continue
    owner='' if name=='_rels/.rels' else name.replace('/_rels/','/')[:-5]
    for rel in E.fromstring(data):
        if rel.get('TargetMode')!='External':
            target=resolve(owner,rel.get('Target'))
            assert target in parts,(name,target)
            relationships+=1

pres=E.fromstring(parts['ppt/presentation.xml'])
rels={r.get('Id'):r.get('Target') for r in E.fromstring(parts['ppt/_rels/presentation.xml.rels'])}
ids=list(pres.find('p:sldIdLst',NS))
paths=[resolve('ppt/presentation.xml',rels[n.get(R+'id')]) for n in ids]
assert len(paths)==34
assert paths==[s['path'] for s in report['slide_order']]
section_ids=[s.get('id') for s in pres.xpath('//*[local-name()="section"]/*[local-name()="sldIdLst"]/*')]
assert section_ids==[n.get('id') for n in ids]
expected_titles=[s['title'] for s in report['slide_order']]
script=(FINAL/'Speaking/Thesis_Defense_Speaking_Script.tex').read_text()
assert re.findall(r'\\section\*\{([^}]+)\}',script)==expected_titles
hidden=[]
notes_count=0
for i,path in enumerate(paths,1):
    root=E.fromstring(parts[path])
    assert (root.get('show')=='0')==(i>25)
    if root.get('show')=='0':hidden.append(i)
    shape_ids=root.xpath('//*[local-name()="cNvPr"]/@id')
    assert len(shape_ids)==len(set(shape_ids)),path
    if i>1:
        foot=root.xpath('//p:sp[p:nvSpPr/p:cNvPr/@name="TextBox 6"]//a:t/text()',namespaces=NS)
        assert ''.join(foot)==report['slide_order'][i-1]['footer'],(i,foot)
    rs=E.fromstring(parts[relpath(path)])
    note=next(r for r in rs if r.get('Type').endswith('/notesSlide'))
    nr=E.fromstring(parts[resolve(path,note.get('Target'))])
    body=nr.xpath('//p:sp[p:nvSpPr/p:nvPr/p:ph/@type="body"]//a:t/text()',namespaces=NS)
    assert body,i
    notes_count+=1
    if i in [20,21,22,23,24,25,28,29,30,31,32]:
        spoken='\n'.join(body).split('[Sources]',1)[0].strip().splitlines()
        section=re.findall(r'\\section\*\{([^}]+)\}(.*?)(?=\\section\*\{|\\end\{document\})',script,re.S)[i-1][1]
        for line in spoken:assert '\\item '+line.replace('%','\\%') in section,(i,line)

# Editable equations and video payloads are byte-identical in original parts.
math=lambda r:[E.tostring(n) for n in r.xpath('//*[local-name()="oMath"]')]
for path in paths:
    if path in before:assert math(E.fromstring(parts[path]))==math(E.fromstring(before[path])),path
videos=[n for n in before if n.endswith(('.mp4','.wmv'))]
assert len(videos)==2
for n in videos:assert parts[n]==before[n],n
catalog=json.loads((ASSETS/'native_equations.json').read_text())
assert len(catalog)==35
for item in catalog:
    root=E.fromstring(parts[paths[item['slide']-1]])
    names=root.xpath('//*[local-name()="cNvPr"]/@name')
    assert item['shape'] in names,item

# Native table values are computed from the same six-condition analysis.
prs=Presentation(FINAL/'Thesis_Defense_gg0_v3.pptx')
table=next(s.table for s in prs.slides[23].shapes if s.has_table)
values=[[c.text for c in row.cells] for row in table.rows]
expected=[[f"{r['k_sigma_Nm']:g}",f"{r['conditioning_mean_deg']:.2f}",f"{r['combined_mean_deg']:.2f}"] for r in metrics['parameter_comparison']]
assert values[1:]==expected
assert metrics['main_cumulative_means_deg']==[analysis['conditions'][i]['cumulative_mean'] for i in [0,1,3,4]]
for row,a,b in [(metrics['parameter_comparison'][0],2,5),(metrics['parameter_comparison'][1],3,4)]:
    assert row['conditioning_mean_deg']==analysis['conditions'][a]['cumulative_mean']
    assert row['combined_mean_deg']==analysis['conditions'][b]['cumulative_mean']
    assert row['conditioning_sample_sd_deg']==analysis['conditions'][a]['cumulative_sd']
    assert row['combined_sample_sd_deg']==analysis['conditions'][b]['cumulative_sd']

# Main pictures use only the selected four conditions. Detailed figures retain
# the original six-condition media payloads from the pre-restructure snapshot.
for n,name in [(21,'nullspace_cumulative_main'),(22,'nullspace_conditioning_main'),(23,'joint_motion_mean_main')]:
    assert parts['ppt/media/'+name+'.png']==(ASSETS/(name+'.png')).read_bytes()
    texts=' '.join(fitz.open(ASSETS/(name+'.pdf'))[0].get_text().split())
    assert '1.5' not in texts
    assert 'Three-trial mean and one sample standard deviation.' in texts
for name,data in before.items():
    if name.startswith('ppt/media/') and name.endswith('.png'):
        assert parts[name]==data,name

pdf=fitz.open(FINAL/'Thesis_Defense_gg0_v3.pdf')
supp=fitz.open(FINAL/'Supplementary_slides.pdf')
oldpdf=fitz.open(HERE/'review/before.pdf')
assert len(pdf)==34 and len(supp)==9
pixels=lambda page:page.get_pixmap(matrix=fitz.Matrix(1,1),alpha=False).samples
for i in range(19):assert pixels(pdf[i])==pixels(oldpdf[i]),i+1
for i in range(9):assert pixels(pdf[i+25])==pixels(supp[i]),i+26
for i,title in enumerate(expected_titles[1:],1):
    assert title in ' '.join(pdf[i].get_text().split()),(i+1,title)
assert all(v in pdf[23].get_text() for row in expected for v in row)
assert all(p.rect==fitz.Rect(0,0,960,540) for p in pdf)

results={
    'slides':34,'visible':25,'hidden_physical_slides':hidden,
    'pdf_pages':len(pdf),'supplementary_pages':len(supp),
    'resolved_internal_relationships':relationships,
    'native_equation_shapes_preserved':len(catalog),'embedded_videos_preserved':len(videos),
    'notes_slides':notes_count,'speaking_title_order_matches':True,
    'native_table_matches_analysis':True,'main_selection_matches_six_condition_analysis':True,
    'original_media_preserved':True,'first_19_pdf_pages_pixel_identical':True,
    'supplementary_pages_pixel_identical_to_full_pdf':True,
    'artifacts_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [FINAL/'Thesis_Defense_gg0_v3.pptx',FINAL/'Thesis_Defense_gg0_v3.pdf',FINAL/'Supplementary_slides.pdf']}}
(HERE/'verification.json').write_text(json.dumps(results,indent=2)+'\n')
print(json.dumps(results,indent=2))
